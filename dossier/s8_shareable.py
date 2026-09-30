"""Stage 8 - the shareable layer: the graph stays, the identity data does not.

Natural persons live in a separate identity layer (``identity/persons.json``: identifier, name and
the other invented personal data). The shareable dossier is built from a copy of the view in which

* every person identifier is replaced by an opaque node id (``N-`` + salted SHA-256, 12 hex digits);
* quotes, line numbers and file names of the source documents are dropped (a quoted line can carry
  a name);

so that the ownership graph keeps its shape - same nodes, same edges, same exact shares - and no
string of the identity layer survives. ``leaks`` checks that on the files actually written.

This separates two layers. It is not an access-control system and does not claim to be one.
"""
from __future__ import annotations

import copy
import hashlib
import re
from pathlib import Path

from docx import Document

from .lib import jsonio
from . import s6_build

_PERSON = re.compile(r"\bP-\d{3}\b")
_DROP = ("quote", "line", "line_end", "source_file", "file")
NAMES = ("dossier_shareable.docx", "provenance_shareable.json")


def opaque(person_id: str, salt: str) -> str:
    return "N-" + hashlib.sha256(f"{salt}:{person_id}".encode("utf-8")).hexdigest()[:12]


def _scrub(obj, salt: str):
    if isinstance(obj, str):
        return _PERSON.sub(lambda m: opaque(m.group(0), salt), obj)
    if isinstance(obj, list):
        return [_scrub(x, salt) for x in obj]
    if isinstance(obj, dict):
        return {_scrub(k, salt): _scrub(v, salt) for k, v in obj.items() if k not in _DROP}
    return obj


def shareable_view(view: dict, salt: str) -> dict:
    if not salt:
        raise ValueError("a salt is required for the shareable layer")
    out = _scrub(copy.deepcopy(view), salt)
    # tables are kept sorted by (opaque) holder so that the output does not leak the original order
    for res in out["fields"].values():
        for holder_table in _tables(res):
            holder_table.sort(key=lambda r: r["holder"])
        if isinstance(res.get("value"), list) and res["field"] == "directors":
            res["value"].sort()
    return out


def _tables(res: dict):
    if res["field"] != "shareholders":
        return
    if isinstance(res.get("value"), list):
        yield res["value"]
    for c in res.get("candidates", []):
        yield c["value"]
    for h in res.get("historical", []):
        yield h["value"]


def graph_of(view: dict) -> dict:
    """Nodes and edges of the ownership chain of a view (internal or shareable)."""
    own = view["ownership"]
    edges = [{"holder": c["holder"], "held": c["held"], "share": c["share"], "source_doc": c["source_doc"],
              "source_date": c["source_date"], "edition": c["edition"]} for c in own["chain"]]
    ids = sorted({view["entity"]["id"]} | {e["holder"] for e in edges} | {e["held"] for e in edges})
    nodes = [{"id": i, "type": "entity" if i.startswith("E-") else "person"} for i in ids]
    return {"entity": view["entity"]["id"], "nodes": nodes,
            "edges": sorted(edges, key=lambda e: (e["held"], e["holder"])),
            "cycles": own["cycles"], "outcome": own["outcome"], "effective": own["effective"],
            "method": own["method"]}


def build_shareable(view: dict, out_dir: Path | str, as_of: str, assumptions: list[dict], salt: str) -> dict:
    sv = shareable_view(view, salt)
    built = s6_build.build(sv, out_dir, as_of, assumptions, identity=None, shareable=True, names=NAMES)
    jsonio.write(Path(out_dir) / "graph.json", graph_of(sv))
    return built


def identity_strings(identity: dict) -> list[str]:
    """Every string of the identity layer: identifiers and each personal datum."""
    out = set()
    for pid, person in identity.items():
        out.add(pid)
        for v in person.values():
            if isinstance(v, str) and v:
                out.add(v)
    return sorted(out)


def leaks(out_dir: Path | str, identity: dict) -> list[str]:
    """Identity-layer strings found in the shareable files of a folder (must be empty)."""
    d = Path(out_dir)
    text = jsonio.read_text(d / NAMES[1]) + jsonio.read_text(d / "graph.json")
    doc = Document(str(d / NAMES[0]))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        parts.extend(c.text for r in t.rows for c in r.cells)
    text += "\n".join(parts)
    found = [s for s in identity_strings(identity) if s in text]
    if _PERSON.search(text):
        found.append("person identifier pattern")
    return sorted(set(found))
