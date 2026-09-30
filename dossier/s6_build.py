"""Stage 6 - the dossier as a build artefact: ``dossier.docx`` and ``provenance.json``.

The dossier is regenerated from the entity view (record -> resolution -> ownership); it is never
edited by hand. The builder refuses to render (``BuildRefused``):

* a figure without ``value``, ``source_doc``, ``source_date`` or ``edition``;
* a figure whose source is not a document of the record;
* a ``[TO CONFIRM]`` entry that carries a value, a discrepancy with fewer than two values
  (what single documents state about a field to confirm is listed apart, as statements, not facts);
* a float anywhere in the view;

and it refuses with ``BuildBlocked`` a cap table that does not sum to exactly the whole.

Two builds from the same view and the same ``as_of`` give the same bytes: fixed core properties,
no wall clock, and a normalised zip container (``lib/zipnorm.py``).
"""
from __future__ import annotations

import datetime as _dt
from fractions import Fraction
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

from . import SYNTHETIC_NOTICE, TO_CONFIRM, __version__
from .lib import jsonio, zipnorm
from .lib.numbers import AMOUNT_CANON, format_amount, parse_frac, share_display

FIELD_ORDER = ["name", "legal_form", "registered_office", "share_capital.resolved", "share_capital.subscribed",
               "share_capital.paid_in", "directors"]
SOURCE = ("source_doc", "source_date", "edition")

H_FACTS = ["Fig.", "Field", "Value", "Nature", "Status", "Source document", "Source date", "Edition"]
H_DISC = ["Fig.", "Field in discrepancy", "Value", "Nature", "Source document", "Source date", "Edition"]
H_TOCONFIRM = ["Field to confirm", "Value", "Why"]
H_UNCONF = ["Fig.", "Field to confirm", "Readable statement (not a fact)", "Nature", "Source document",
            "Source date", "Edition"]
H_CAP = ["Fig.", "Holder", "Share (exact)", "Source document", "Source date", "Edition"]
H_CHAIN = ["Fig.", "Holder", "Held entity", "Share (exact)", "Source document", "Source date", "Edition"]
H_EFF = ["Fig.", "Natural person", "Effective share (exact)", "Method", "Derived from"]
H_HIST = ["Fig.", "Field (history)", "Historical value", "Superseded by", "Source document", "Source date", "Edition"]
H_DOCS = ["Document", "Type", "Date", "Edition", "SHA-256"]
H_IDENT = ["Person", "Name (restricted identity layer)"]
H_ASSUME = ["Parameter", "Value", "Status", "Note"]


class BuildRefused(RuntimeError):
    """The view would render a figure without a source, or is malformed. The run is FAILED."""


class BuildBlocked(RuntimeError):
    """The cap table does not sum to the whole. The run is BLOCKED and nothing is published."""


def field_order(fields: dict) -> list[str]:
    known = [f for f in FIELD_ORDER if f in fields]
    return known + sorted(f for f in fields if f not in FIELD_ORDER and f != "shareholders")


def display(fld: str, value) -> str:
    """How a value is written in the dossier."""
    if fld == "shareholders":
        return "; ".join(f"{r['holder']} {share_display(r['share'])}" for r in value)
    if fld == "directors":
        return ", ".join(value)
    if fld.startswith(("share_capital.", "fin.")):
        return format_amount(value)
    if fld in ("share", "effective"):
        return share_display(value)
    return str(value)


def _no_float(obj, where: str = "view") -> None:
    if isinstance(obj, float):
        raise BuildRefused(f"float found at {where}: values are decimal or fraction strings")
    if isinstance(obj, dict):
        for k, v in obj.items():
            _no_float(v, f"{where}.{k}")
    elif isinstance(obj, (list, tuple)):
        for i, v in enumerate(obj):
            _no_float(v, f"{where}[{i}]")


def _figure(section: str, fld: str, status: str, value, src: dict, **extra) -> dict:
    fig = {"section": section, "field": fld, "status": status, "value": value,
           "display": None if value is None else display(extra.pop("display_as", fld), value)}
    for k in ("nature", "source_doc", "source_file", "source_date", "edition", "line", "line_end", "quote"):
        if k in src:
            fig[k] = src[k]
    fig.update(extra)
    return fig


def model_from_view(view: dict) -> dict:
    """Every figure the dossier will show, in a fixed order, numbered F-001, F-002, ..."""
    _no_float(view)
    figures: list[dict] = []
    to_confirm: list[dict] = []
    notes: list[str] = []
    fields = view["fields"]
    for fld in field_order(fields) + (["shareholders"] if "shareholders" in fields else []):
        res = fields[fld]
        status = res.get("status")
        if status == "STATED":
            if "value" not in res or not res.get("sources"):
                raise BuildRefused(f"{fld}: STATED without a value or without a source")
            first, others = res["sources"][0], res["sources"][1:]
            also = [{k: s.get(k) for k in SOURCE} for s in others]
            if fld == "shareholders":
                for row in res["value"]:
                    figures.append(_figure("cap_table", fld, "STATED", row["share"], first, display_as="share",
                                           holder=row["holder"], also_stated_by=also))
            else:
                figures.append(_figure("facts", fld, "STATED", res["value"], first, also_stated_by=also))
        elif status == "DISCREPANCY":
            cands = res.get("candidates") or []
            if len(cands) < 2 or "value" in res:
                raise BuildRefused(f"{fld}: a discrepancy shows every value side by side, never one")
            for n, cand in enumerate(cands, 1):
                if not cand.get("sources"):
                    raise BuildRefused(f"{fld}: discrepancy value without a source")
                for s in cand["sources"]:
                    figures.append(_figure("discrepancies", fld, "DISCREPANCY", cand["value"], s, candidate=n))
        elif status == "TO_CONFIRM":
            if "value" in res or res.get("candidates"):
                raise BuildRefused(f"{fld}: a field to confirm carries no value")
            to_confirm.append({"field": fld, "marker": TO_CONFIRM, "why": res.get("reason", "")})
            for n, cand in enumerate(res.get("readable", []), 1):
                for s in cand["sources"]:
                    figures.append(_figure("unconfirmed", fld, "UNCONFIRMED", cand["value"], s, candidate=n))
        else:
            raise BuildRefused(f"{fld}: unknown status {status!r}")
        for u in res.get("unreadable", []):
            if status != "TO_CONFIRM":
                notes.append(f"{fld}: {u['source_doc']} ({u['edition']}, {u['source_date']}) could not be read "
                             f"for this field - {u['reason']}.")

    own = view["ownership"]
    if own["outcome"] == "BLOCKED":
        raise BuildBlocked(own["reason"])
    for c in own["sum_checks"]:
        if not c["whole"]:
            raise BuildBlocked(f"holders' table of {c['source_doc']} sums to {c['sum']}, not to the whole")
    cap = [f for f in figures if f["section"] == "cap_table"]
    if cap and sum((parse_frac(f["value"]) for f in cap), Fraction(0)) != 1:
        raise BuildBlocked("the cap table to render does not sum to the whole")
    target = view["entity"]["id"]
    for edge in own["chain"]:
        if edge["held"] != target:
            figures.append(_figure("chain", "shareholders", "STATED", edge["share"], edge, display_as="share",
                                   holder=edge["holder"], held=edge["held"]))

    for fld in field_order(fields) + (["shareholders"] if "shareholders" in fields else []):
        for h in fields[fld].get("historical", []):
            by = h["superseded_by"]
            figures.append(_figure("history", fld, "HISTORICAL", h["value"], h,
                                   superseded_by=f"{by['source_doc']} ({by['edition']}, {by['source_date']})"))
    for p in view.get("previous_statements", []):
        figures.append(_figure("history", p["field"].replace(".previous", ""), "HISTORICAL", p["value"], p,
                               display_as="share_capital." if p["field"].startswith("share_capital") else "text",
                               superseded_by="stated as the previous value by the same document"))

    for i, f in enumerate(figures, 1):
        f["id"] = f"F-{i:03d}"
    docs = {d["doc_id"] for d in view["documents"]} | set(view.get("upstream_documents", {}))
    for f in figures:
        for k in ("value",) + SOURCE:
            if f.get(k) in (None, "", []):
                raise BuildRefused(f"figure {f['id']} ({f['field']}) has no {k}: a figure without a source is not rendered")
        if f["source_doc"] not in docs:
            raise BuildRefused(f"figure {f['id']} cites {f['source_doc']}, which is not a document of the record")
        if f["field"].startswith(("share_capital.", "fin.")) and not AMOUNT_CANON.match(str(f["value"])):
            raise BuildRefused(f"figure {f['id']}: amount is not a canonical decimal string")

    derived: list[dict] = []
    if own["outcome"] == "RESOLVED":
        edge_ids = [f["id"] for f in figures if f["section"] in ("cap_table", "chain")]
        if not edge_ids:
            raise BuildRefused("effective holdings without any sourced edge")
        for n, (person, share) in enumerate(sorted(own["effective"].items()), 1):
            derived.append({"id": f"D-{n:03d}", "section": "effective", "person": person, "value": share,
                            "display": share_display(share), "method": own["method"], "derived_from": edge_ids})
    elif own["outcome"] == "TO_CONFIRM":
        to_confirm.append({"field": "effective_holdings", "marker": TO_CONFIRM, "why": own["reason"]})
    else:
        raise BuildRefused(f"unknown ownership outcome {own['outcome']!r}")
    return {"figures": figures, "derived": derived, "to_confirm": to_confirm, "notes": notes}


# ----------------------------------------------------------------------------------------------
# DOCX

def _table(doc, header: list[str], rows: list[list[str]], red_col: int | None = None) -> None:
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
    for r, row in enumerate(rows, 1):
        cells = t.rows[r].cells
        for i, v in enumerate(row):
            if i == red_col:
                cells[i].text = ""
                run = cells[i].paragraphs[0].add_run(str(v))
                run.bold, run.font.color.rgb = True, RGBColor(0xC0, 0x00, 0x00)
            else:
                cells[i].text = str(v)


def _src_cells(f: dict) -> list[str]:
    return [f["source_doc"], f["source_date"], f["edition"]]


def render_docx(path: Path, view: dict, model: dict, as_of: str, rules_assumptions: list[dict],
                identity: dict | None = None, shareable: bool = False) -> None:
    ent = view["entity"]
    doc = Document()
    doc.styles["Normal"].font.size = Pt(9)
    section = doc.sections[0]
    section.header.paragraphs[0].text = SYNTHETIC_NOTICE
    section.footer.paragraphs[0].text = f"{SYNTHETIC_NOTICE} Data as of {as_of}."
    layer = "shareable layer (no identity data)" if shareable else "internal layer"
    doc.add_heading(f"Entity dossier {ent['id']} - {layer}", level=0)
    doc.add_paragraph(f"Test registry number {ent['registry_no']}. Data as of {as_of}. "
                      f"Rebuilt from the entity record; this file is a build artefact and is never edited by hand.")
    doc.add_paragraph(SYNTHETIC_NOTICE)

    figs = model["figures"]
    doc.add_heading("1. Entity facts - every figure with its source, date and edition", level=1)
    _table(doc, H_FACTS, [[f["id"], f["field"], f["display"], f.get("nature", ""), f["status"]] + _src_cells(f)
                          for f in figs if f["section"] == "facts"])
    also = [f for f in figs if f["section"] == "facts" and f.get("also_stated_by")]
    for f in also:
        doc.add_paragraph(f"{f['id']} is also stated by: "
                          + "; ".join(f"{s['source_doc']} ({s['edition']}, {s['source_date']})"
                                      for s in f["also_stated_by"]) + ".", style="List Bullet")

    doc.add_heading("2. Discrepancies - side by side, not adjudicated", level=1)
    disc = [f for f in figs if f["section"] == "discrepancies"]
    if disc:
        doc.add_paragraph("The sources below disagree. Both values are shown with their sources; "
                          "which one prevails is for a human to decide.")
        _table(doc, H_DISC, [[f["id"], f["field"], f["display"], f.get("nature", "")] + _src_cells(f) for f in disc])
    else:
        doc.add_paragraph("None found between the current sources.")

    doc.add_heading("3. To confirm", level=1)
    if model["to_confirm"]:
        _table(doc, H_TOCONFIRM, [[t["field"], t["marker"], t["why"]] for t in model["to_confirm"]], red_col=1)
    else:
        doc.add_paragraph("Nothing to confirm.")
    unconf = [f for f in figs if f["section"] == "unconfirmed"]
    if unconf:
        doc.add_paragraph("What could be read for the fields above. These are statements of single documents, "
                          "not facts of this dossier: a current source could not be read, or an unread document "
                          "may change them.")
        _table(doc, H_UNCONF, [[f["id"], f["field"], f["display"], f.get("nature", "")] + _src_cells(f)
                               for f in unconf])

    doc.add_heading("4. Ownership", level=1)
    own = view["ownership"]
    cap = [f for f in figs if f["section"] == "cap_table"]
    if cap:
        total = sum((parse_frac(f["value"]) for f in cap), Fraction(0))
        rows = [[f["id"], f["holder"], f["display"]] + _src_cells(f) for f in cap]
        rows.append(["", "Sum", share_display(f"{total.numerator}/{total.denominator}"), "", "", ""])
        _table(doc, H_CAP, rows)
    else:
        doc.add_paragraph("No single holders' table can be shown: see sections 2 and 3.")
    chain = [f for f in figs if f["section"] == "chain"]
    if chain:
        doc.add_paragraph("Chain above the holders (tables of the holder entities):")
        _table(doc, H_CHAIN, [[f["id"], f["holder"], f["held"], f["display"]] + _src_cells(f) for f in chain])
    for cyc in own["cycles"]:
        doc.add_paragraph("Cross-holding found: " + " <-> ".join(cyc) + ".", style="List Bullet")
    if model["derived"]:
        doc.add_paragraph("Effective holding of natural persons (derived, exact):")
        _table(doc, H_EFF, [[d["id"], d["person"], d["display"], d["method"],
                             f"{d['derived_from'][0]} to {d['derived_from'][-1]}"] for d in model["derived"]])

    doc.add_heading("5. History - superseded values, each bound to its edition", level=1)
    hist = [f for f in figs if f["section"] == "history"]
    if hist:
        _table(doc, H_HIST, [[f["id"], f["field"], f["display"], f["superseded_by"]] + _src_cells(f) for f in hist])
    else:
        doc.add_paragraph("No superseded value.")

    doc.add_heading("6. Notes and warnings", level=1)
    w = view["warnings"]
    lines = list(model["notes"])
    for d in w["filename_divergences"]:
        lines.append(f"{d['doc_id']}: the file name suggests {d['aspect']} {d['filename_says']}, the content says "
                     f"{d['content_says']}. The content is used; the file name is never a source.")
    for d in w["unclassified_documents"]:
        lines.append(f"A document was not read ({d.get('doc_id') or 'no id'}): {d['reason']}.")
    for d in w["ignored_after_as_of"]:
        lines.append(f"{d} is dated after {as_of} and was not considered.")
    for line in lines or ["None."]:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("7. Source documents", level=1)
    all_docs = list(view["documents"]) + [d for _, d in sorted(view.get("upstream_documents", {}).items())]
    _table(doc, H_DOCS, [[d["doc_id"], d["kind"], d["date"], d["edition"], d["sha256"]] for d in all_docs])

    if identity and not shareable:
        doc.add_heading("8. Natural persons - restricted identity layer", level=1)
        used = sorted({p for f in figs for p in _persons_of(f)} | {d["person"] for d in model["derived"]})
        _table(doc, H_IDENT, [[p, identity.get(p, {}).get("name", TO_CONFIRM)] for p in used])

    doc.add_heading("Assumptions - every one is [TO CONFIRM with legal]", level=1)
    doc.add_paragraph("This dossier is not legal advice. Each assumption below is a parameter of a rule file.")
    _table(doc, H_ASSUME, [[f"{a['file']} - {a['parameter']}", str(a["value"]), a["status"], a["note"]]
                           for a in rules_assumptions])

    cp = doc.core_properties
    t0 = _dt.datetime.fromisoformat(as_of)
    cp.author = "Ezio Cardone (synthetic alumnus, AI agent)"
    cp.last_modified_by = f"dossier.s6_build {__version__}"
    cp.title = f"Entity dossier {ent['id']} (SYNTHETIC)"
    cp.subject = "Legal-entity dossier - synthetic test data"
    cp.comments = "Build artefact - do not edit by hand. Not legal advice."
    cp.keywords = "SYNTHETIC"
    cp.created = cp.modified = cp.last_printed = t0
    cp.revision = 1
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    zipnorm.normalize(path, t0)


def _persons_of(f: dict) -> list[str]:
    out = []
    if f.get("holder", "").startswith("P-"):
        out.append(f["holder"])
    if f["field"] == "directors" and isinstance(f["value"], list):
        out.extend(f["value"])
    if f["field"] == "shareholders" and isinstance(f["value"], list):
        out.extend(r["holder"] for r in f["value"] if r["holder"].startswith("P-"))
    return out


def build(view: dict, out_dir: Path | str, as_of: str, assumptions: list[dict], identity: dict | None = None,
          shareable: bool = False, names: tuple[str, str] = ("dossier.docx", "provenance.json")) -> dict:
    """Write the DOCX and the provenance file. Raises BuildRefused / BuildBlocked before writing anything."""
    model = model_from_view(view)
    out = Path(out_dir)
    provenance = {
        "schema": "provenance/2.0.0",
        "generator": f"dossier {__version__}",
        "layer": "shareable" if shareable else "internal",
        "entity": view["entity"],
        "as_of": as_of,
        "documents": view["documents"],
        "upstream_documents": view.get("upstream_documents", {}),
        "figures": model["figures"],
        "derived": model["derived"],
        "to_confirm": model["to_confirm"],
        "notes": model["notes"],
        "field_status": {f: r["status"] for f, r in sorted(view["fields"].items())},
        "ownership": {k: view["ownership"][k] for k in ("outcome", "rule", "cycle_policy", "cycles", "reason",
                                                        "method", "sum_checks", "unresolved_nodes")},
        "superseded": view["superseded"],
        "warnings": view["warnings"],
        "legal_assumptions": assumptions,
    }
    docx_path = out / names[0]
    render_docx(docx_path, view, model, as_of, assumptions, identity=identity, shareable=shareable)
    jsonio.write(out / names[1], provenance)
    return {"docx": docx_path, "provenance": out / names[1], "model": model,
            "docx_sha256": jsonio.sha256_file(docx_path), "provenance_sha256": jsonio.sha256_file(out / names[1])}
